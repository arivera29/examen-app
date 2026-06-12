import re
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from app.application.exam_report_export import (
    ALT_ROW_FILL,
    CORRECT_FONT,
    HEADER_FILL,
    HEADER_FONT,
    INCORRECT_FONT,
    SECTION_FONT,
    SUBTITLE_FONT,
    THIN_BORDER,
    TITLE_FONT,
    VALUE_FONT,
    _safe_filename,
)

QUESTION_TYPE_LABELS = {
    "single_choice": "Selección única",
    "multiple_choice": "Selección múltiple",
    "open": "Abierta",
    "true_false": "Verdadero/Falso",
}


def build_attempts_answers_filename(exam_title: str, invitee_email: str | None = None) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    base = f"respuestas-{_safe_filename(exam_title)}"
    if invitee_email:
        email_part = _safe_filename(invitee_email.split("@")[0]) or "invitado"
        return f"{base}-{email_part}-{stamp}.xlsx"
    return f"{base}-{stamp}.xlsx"


def build_attempts_answers_workbook(exam_title: str, attempt_reports: list[dict]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen invitados"
    answers_sheet = workbook.create_sheet("Respuestas detalladas")

    generated_at = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    _build_summary_sheet(summary_sheet, exam_title, generated_at, attempt_reports)
    _build_answers_sheet(answers_sheet, attempt_reports)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_summary_sheet(sheet, exam_title: str, generated_at: str, attempt_reports: list[dict]) -> None:
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:I1")
    title_cell = sheet["A1"]
    title_cell.value = f"Respuestas del examen: {exam_title}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A2:I2")
    meta_cell = sheet["A2"]
    meta_cell.value = (
        f"Generado: {generated_at}  |  Intentos exportados: {len(attempt_reports)}  |  "
        f"Invitados: {len({report['invitee_email'] for report in attempt_reports})}"
    )
    meta_cell.font = SUBTITLE_FONT

    headers = [
        "Nombre",
        "Correo",
        "Intento",
        "Calificación",
        "Puntaje máximo",
        "Porcentaje",
        "Acertadas",
        "No acertadas",
        "Sin responder",
    ]
    header_row = 4
    sheet.merge_cells(start_row=header_row - 1, start_column=1, end_row=header_row - 1, end_column=len(headers))
    section_cell = sheet.cell(row=header_row - 1, column=1, value="Resumen por intento")
    section_cell.font = SECTION_FONT

    _write_header_row(sheet, header_row, headers)

    for row_index, report in enumerate(attempt_reports, start=header_row + 1):
        answers = report.get("answers") or []
        correct = sum(1 for answer in answers if answer.get("is_correct") is True)
        incorrect = sum(1 for answer in answers if answer.get("is_correct") is False)
        unanswered = sum(1 for answer in answers if not answer.get("answered"))
        score = report.get("score")
        total = report.get("total_score") or 0
        percentage = round((score / total * 100), 1) if score is not None and total else None

        values = [
            report.get("invitee_full_name") or "-",
            report.get("invitee_email") or "-",
            report.get("attempt_number") or 1,
            score,
            total,
            percentage,
            correct,
            incorrect,
            unanswered,
        ]

        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = THIN_BORDER
            cell.font = VALUE_FONT
            cell.alignment = Alignment(
                horizontal="left" if col_index <= 2 else "center",
                vertical="center",
            )
            if row_index % 2 == 0:
                cell.fill = ALT_ROW_FILL
            if col_index == 7:
                cell.font = CORRECT_FONT
            elif col_index == 8:
                cell.font = INCORRECT_FONT
            elif col_index == 6 and isinstance(value, (int, float)):
                cell.number_format = "0.0"

    column_widths = [28, 30, 10, 12, 14, 12, 12, 14, 14]
    for index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate


def _build_answers_sheet(sheet, attempt_reports: list[dict]) -> None:
    sheet.sheet_view.showGridLines = False
    headers = [
        "Nombre",
        "Correo",
        "Intento",
        "#",
        "Pregunta",
        "Tipo",
        "Respuesta del invitado",
        "Respuesta correcta",
        "Resultado",
        "Puntos",
        "Máx.",
    ]

    header_row = 3
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(row=1, column=1, value="Detalle de preguntas y respuestas")
    title_cell.font = SECTION_FONT

    _write_header_row(sheet, header_row, headers)

    row_index = header_row + 1
    for report in attempt_reports:
        invitee_name = report.get("invitee_full_name") or "-"
        invitee_email = report.get("invitee_email") or "-"
        attempt_number = report.get("attempt_number") or 1

        for answer in report.get("answers") or []:
            result = _result_label(answer)
            values = [
                invitee_name,
                invitee_email,
                attempt_number,
                answer.get("order"),
                answer.get("question_text"),
                QUESTION_TYPE_LABELS.get(answer.get("question_type"), answer.get("question_type")),
                _answer_text(answer),
                "; ".join(answer.get("correct_options") or []) or "-",
                result,
                answer.get("points_earned", 0),
                answer.get("max_points", 0),
            ]

            for col_index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                cell.border = THIN_BORDER
                cell.font = VALUE_FONT
                cell.alignment = Alignment(
                    horizontal="center" if col_index in {3, 4, 10, 11} else "left",
                    vertical="top",
                    wrap_text=col_index in {5, 7, 8},
                )
                if row_index % 2 == 0:
                    cell.fill = ALT_ROW_FILL
                if col_index == 9:
                    if result == "Correcta":
                        cell.font = CORRECT_FONT
                    elif result == "Incorrecta":
                        cell.font = INCORRECT_FONT

            row_index += 1

    column_widths = [24, 28, 10, 6, 42, 16, 32, 28, 16, 10, 8]
    for index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate


def _write_header_row(sheet, row: int, headers: list[str]) -> None:
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[row].height = 28


def _answer_text(answer: dict) -> str:
    open_text = (answer.get("open_text") or "").strip()
    if open_text:
        return open_text
    selected = answer.get("selected_options") or []
    if selected:
        return "; ".join(selected)
    return "Sin respuesta"


def _result_label(answer: dict) -> str:
    if not answer.get("answered"):
        return "Sin respuesta"
    if answer.get("is_correct") is None:
        return "Pendiente de revisión"
    return "Correcta" if answer.get("is_correct") else "Incorrecta"
