import re
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.domain.enums import QuestionType

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="666666")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E79")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="333333")
VALUE_FONT = Font(name="Calibri", size=10, color="333333")
CORRECT_FONT = Font(name="Calibri", size=10, bold=True, color="2E7D32")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ALT_ROW_FILL = PatternFill(start_color="F5F8FC", end_color="F5F8FC", fill_type="solid")
SUMMARY_LABEL_FILL = PatternFill(start_color="E8EEF4", end_color="E8EEF4", fill_type="solid")
SUMMARY_VALUE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

QUESTION_TYPE_LABELS = {
    QuestionType.SINGLE_CHOICE.value: "Selección única",
    QuestionType.MULTIPLE_CHOICE.value: "Selección múltiple",
    QuestionType.OPEN.value: "Abierta",
    QuestionType.TRUE_FALSE.value: "Verdadero/Falso",
}


def _safe_filename(title: str) -> str:
    safe = re.sub(r"[^\w\-]+", "-", title).strip("-")
    return safe or "banco"


def build_export_filename(payload: dict) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"banco-preguntas-{_safe_filename(payload['bank_name'])}-{stamp}.xlsx"


def build_question_bank_workbook(payload: dict) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    questions_sheet = workbook.create_sheet("Preguntas")

    _build_summary_sheet(summary_sheet, payload)
    _build_questions_sheet(questions_sheet, payload)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_summary_sheet(sheet, payload: dict) -> None:
    sheet.sheet_view.showGridLines = False
    generated_at = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    questions = payload["questions"]
    total = payload["total_questions"]

    sheet.merge_cells("A1:F1")
    title_cell = sheet["A1"]
    title_cell.value = f"Banco de preguntas: {payload['bank_name']}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A2:F2")
    meta_cell = sheet["A2"]
    meta_cell.value = f"Generado: {generated_at}  |  Total de preguntas: {total}"
    meta_cell.font = SUBTITLE_FONT

    description = (payload.get("bank_description") or "").strip()
    if description:
        sheet.merge_cells("A3:F3")
        desc_cell = sheet["A3"]
        desc_cell.value = f"Descripción: {description}"
        desc_cell.font = SUBTITLE_FONT
        desc_cell.alignment = Alignment(wrap_text=True)

    section_row = 5 if description else 4
    sheet.merge_cells(start_row=section_row, start_column=1, end_row=section_row, end_column=6)
    section_cell = sheet.cell(row=section_row, column=1, value="Resumen por tipo de pregunta")
    section_cell.font = SECTION_FONT

    type_counts = Counter(
        QUESTION_TYPE_LABELS.get(q["question_type"], q["question_type"]) for q in questions
    )
    header_row = section_row + 1
    for col_index, header in enumerate(["Tipo", "Cantidad"], start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_offset, (label, count) in enumerate(sorted(type_counts.items()), start=1):
        row = header_row + row_offset
        label_cell = sheet.cell(row=row, column=1, value=label)
        value_cell = sheet.cell(row=row, column=2, value=count)
        for cell in (label_cell, value_cell):
            cell.border = THIN_BORDER
            cell.font = VALUE_FONT
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center")
            if row % 2 == 0:
                cell.fill = ALT_ROW_FILL

    topic_section_row = header_row + len(type_counts) + 3
    sheet.merge_cells(start_row=topic_section_row, start_column=1, end_row=topic_section_row, end_column=6)
    topic_title = sheet.cell(row=topic_section_row, column=1, value="Resumen por tema")
    topic_title.font = SECTION_FONT

    topic_counts = Counter((q.get("topic_name") or "Sin tema") for q in questions)
    topic_header_row = topic_section_row + 1
    for col_index, header in enumerate(["Tema", "Cantidad"], start=1):
        cell = sheet.cell(row=topic_header_row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_offset, (label, count) in enumerate(sorted(topic_counts.items()), start=1):
        row = topic_header_row + row_offset
        label_cell = sheet.cell(row=row, column=1, value=label)
        value_cell = sheet.cell(row=row, column=2, value=count)
        for cell in (label_cell, value_cell):
            cell.border = THIN_BORDER
            cell.font = VALUE_FONT
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "center")
            if row % 2 == 0:
                cell.fill = ALT_ROW_FILL

    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 14


def _build_questions_sheet(sheet, payload: dict) -> None:
    sheet.sheet_view.showGridLines = False
    headers = [
        "#",
        "Enunciado",
        "Tipo",
        "Tema",
        "Tiempo (s)",
        "Imagen",
        "Respuestas correctas",
        "Todas las opciones",
    ]

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(row=1, column=1, value="Detalle de preguntas y respuestas")
    title_cell.font = SECTION_FONT
    title_cell.alignment = Alignment(horizontal="left")

    header_row = 3
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    sheet.row_dimensions[header_row].height = 28

    for index, question in enumerate(payload["questions"], start=1):
        row = header_row + index
        correct_text, options_text = _format_answers(question)
        values = [
            index,
            question["text"],
            QUESTION_TYPE_LABELS.get(question["question_type"], question["question_type"]),
            question.get("topic_name") or "Sin tema",
            question["time_seconds"],
            "Sí" if question.get("image_url") else "No",
            correct_text,
            options_text,
        ]

        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col_index, value=value)
            cell.border = THIN_BORDER
            cell.font = CORRECT_FONT if col_index == 7 and value not in {"-", "Evaluación manual"} else VALUE_FONT
            cell.alignment = Alignment(
                horizontal="center" if col_index in {1, 4, 5, 6} else "left",
                vertical="top",
                wrap_text=col_index in {2, 7, 8},
            )
            if row % 2 == 0:
                cell.fill = ALT_ROW_FILL

    column_widths = [6, 48, 18, 22, 12, 10, 32, 42]
    for index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate


def _format_answers(question: dict) -> tuple[str, str]:
    question_type = question["question_type"]
    options = question.get("options") or []

    if question_type == QuestionType.OPEN.value:
        return "Evaluación manual", "-"

    if not options:
        return "-", "-"

    labels = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    formatted_options = []
    correct_answers = []

    for index, option in enumerate(options):
        label = labels[index] if index < len(labels) else str(index + 1)
        marker = " ✓" if option.get("is_correct") else ""
        line = f"{label}) {option['text']}{marker}"
        formatted_options.append(line)
        if option.get("is_correct"):
            correct_answers.append(option["text"])

    return (
        "; ".join(correct_answers) if correct_answers else "-",
        "\n".join(formatted_options),
    )
