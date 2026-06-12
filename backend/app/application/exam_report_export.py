import re
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="666666")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1F4E79")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="333333")
VALUE_FONT = Font(name="Calibri", size=10, color="333333")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ALT_ROW_FILL = PatternFill(start_color="F5F8FC", end_color="F5F8FC", fill_type="solid")
HIGH_FRAUD_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
CORRECT_FONT = Font(name="Calibri", size=10, color="2E7D32")
INCORRECT_FONT = Font(name="Calibri", size=10, color="C62828")

STATUS_LABELS = {
    "submitted": "Enviado",
    "in_progress": "En progreso",
    "timed_out": "Tiempo agotado",
    "not_started": "No iniciado",
}


def _status_label(status: str) -> str:
    return STATUS_LABELS.get(status, status)


def _safe_filename(title: str) -> str:
    safe = re.sub(r"[^\w\-]+", "-", title).strip("-")
    return safe or "examen"


def build_export_filename(report: dict) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"informe-{_safe_filename(report['exam_title'])}-{stamp}.xlsx"


def build_exam_report_workbook(report: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Informe"
    sheet.sheet_view.showGridLines = False

    generated_at = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    summary = report["summary"]
    rows = report["individual_reports"]

    sheet.merge_cells("A1:N1")
    title_cell = sheet["A1"]
    title_cell.value = f"Informe de examen: {report['exam_title']}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    sheet.merge_cells("A2:N2")
    meta_cell = sheet["A2"]
    meta_cell.value = (
        f"Generado: {generated_at}  |  Preguntas: {report.get('question_count', '-')}  |  "
        f"Puntaje máximo: {report['total_score']}"
    )
    meta_cell.font = SUBTITLE_FONT
    meta_cell.alignment = Alignment(horizontal="left")

    sheet.merge_cells("A4:N4")
    section_cell = sheet["A4"]
    section_cell.value = "Resumen general"
    section_cell.font = SECTION_FONT

    summary_items = [
        ("Invitados", summary["total_invitees"]),
        ("Completados", summary["completed"]),
        ("Promedio", round(summary["average_score"], 1)),
        ("Calificación máxima", round(summary["highest_score"], 1)),
        ("Calificación mínima", round(summary["lowest_score"], 1)),
        ("Prom. acertadas", round(summary.get("average_correct_answers") or 0, 1)),
        ("Prom. no acertadas", round(summary.get("average_incorrect_answers") or 0, 1)),
    ]

    summary_start_row = 5
    for index, (label, value) in enumerate(summary_items):
        row = summary_start_row + (index // 4)
        col = (index % 4) * 3 + 1
        label_cell = sheet.cell(row=row, column=col, value=label)
        value_cell = sheet.cell(row=row, column=col + 1, value=value)
        label_cell.font = LABEL_FONT
        value_cell.font = VALUE_FONT
        label_cell.fill = PatternFill(start_color="E8EEF4", end_color="E8EEF4", fill_type="solid")
        value_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        label_cell.border = THIN_BORDER
        value_cell.border = THIN_BORDER
        sheet.merge_cells(start_row=row, start_column=col + 1, end_row=row, end_column=col + 2)
        value_cell.alignment = Alignment(horizontal="center")

    header_row = 9
    headers = [
        "Nombre",
        "Correo",
        "Intento final",
        "Intentos realizados",
        "Calificación",
        "Puntaje máximo",
        "Porcentaje",
        "Acertadas",
        "No acertadas",
        "Estado",
        "Fraude (%)",
        "Eventos supervisión",
        "Inicio",
        "Envío",
    ]

    sheet.merge_cells(start_row=header_row - 1, start_column=1, end_row=header_row - 1, end_column=len(headers))
    detail_title = sheet.cell(row=header_row - 1, column=1, value="Calificaciones individuales")
    detail_title.font = SECTION_FONT
    detail_title.alignment = Alignment(horizontal="left")

    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    sheet.row_dimensions[header_row].height = 28

    for row_index, item in enumerate(rows, start=header_row + 1):
        values = [
            item.get("invitee_full_name") or "-",
            item.get("invitee_email") or "-",
            item.get("attempt_number"),
            item.get("attempts_used") or 0,
            item.get("score"),
            item.get("total_score"),
            item.get("percentage"),
            item.get("correct_answers_count") or 0,
            item.get("incorrect_answers_count") or 0,
            _status_label(str(item.get("status") or "")),
            round((item.get("fraud_score") or 0) * 100, 1),
            item.get("proctoring_events_count") or 0,
            _format_datetime(item.get("started_at")),
            _format_datetime(item.get("submitted_at")),
        ]

        use_alt = row_index % 2 == 0
        high_fraud = (item.get("fraud_score") or 0) > 0.5

        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = THIN_BORDER
            cell.font = VALUE_FONT
            cell.alignment = Alignment(horizontal="left" if col_index <= 2 else "center", vertical="center")

            if high_fraud:
                cell.fill = HIGH_FRAUD_FILL
            elif use_alt:
                cell.fill = ALT_ROW_FILL

            if col_index == 8:
                cell.font = CORRECT_FONT
            elif col_index == 9:
                cell.font = INCORRECT_FONT
            elif col_index == 7 and isinstance(value, (int, float)):
                cell.number_format = "0.0%"
                cell.value = value / 100 if value else 0
            elif col_index in {5, 6} and isinstance(value, (int, float)):
                cell.number_format = "0.0"
            elif col_index == 11 and isinstance(value, (int, float)):
                cell.number_format = "0.0"

    column_widths = [28, 30, 12, 16, 12, 14, 12, 12, 14, 14, 12, 18, 18, 18]
    for index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _format_datetime(value: str | None) -> str:
    if not value:
        return "-"
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return parsed.strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return value
